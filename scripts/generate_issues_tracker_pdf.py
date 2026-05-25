#!/usr/bin/env python
"""Generate a QuickBridge issue tracker PDF from one canonical source."""

from __future__ import annotations

import argparse
import re
from dataclasses import dataclass
from pathlib import Path
from typing import Iterable

from reportlab.lib import colors
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle


ISSUE_RE = re.compile(
    r"\b(?P<key>(?:ISSUE|NEW)-\d+)\b\s*(?:\||-)?\s*(?P<title>.*?)\s*(?:\||-)\s*(?P<status>PARTIAL|OPEN|Partially Addressed|Fully Open)\b",
    re.IGNORECASE,
)


@dataclass(frozen=True)
class Issue:
    key: str
    title: str
    status: str


def read_source(path: Path) -> str:
    if path.suffix.lower() == ".docx":
        from docx import Document

        document = Document(str(path))
        parts: list[str] = []
        parts.extend(paragraph.text for paragraph in document.paragraphs if paragraph.text.strip())
        for table in document.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text.strip() for cell in row.cells))
        return "\n".join(parts)
    return path.read_text(encoding="utf-8")


def parse_issues(source_text: str) -> list[Issue]:
    issues: list[Issue] = []
    seen: set[str] = set()
    for line in source_text.splitlines():
        match = ISSUE_RE.search(" ".join(line.split()))
        if not match:
            continue
        key = match.group("key").upper()
        if key in seen:
            continue
        raw_status = match.group("status").lower()
        status = "PARTIAL" if "partial" in raw_status else "OPEN"
        title = match.group("title").strip(" |-\t")
        issues.append(Issue(key=key, title=title, status=status))
        seen.add(key)
    return issues


def status_counts(issues: Iterable[Issue]) -> dict[str, int]:
    counts = {"PARTIAL": 0, "OPEN": 0}
    for issue in issues:
        counts[issue.status] = counts.get(issue.status, 0) + 1
    counts["TOTAL"] = counts.get("PARTIAL", 0) + counts.get("OPEN", 0)
    return counts


def build_pdf(source: Path, output: Path) -> None:
    issues = parse_issues(read_source(source))
    counts = status_counts(issues)
    if not issues:
        raise SystemExit(f"No issue rows found in {source}")

    doc = SimpleDocTemplate(str(output), pagesize=letter, rightMargin=54, leftMargin=54, topMargin=54, bottomMargin=54)
    styles = getSampleStyleSheet()
    story = [
        Paragraph("QuickBridge Web Connector Issues Tracker 3.0", styles["Title"]),
        Paragraph(f"Source: {source.name}", styles["Normal"]),
        Spacer(1, 18),
    ]

    metric_table = Table(
        [
            ["Status", "Count"],
            ["Partially Addressed", counts["PARTIAL"]],
            ["Fully Open", counts["OPEN"]],
            ["Total Pending", counts["TOTAL"]],
        ],
        colWidths=[260, 100],
    )
    metric_table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1f2937")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#cbd5e1")),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("ALIGN", (1, 1), (1, -1), "CENTER"),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 8),
                ("TOPPADDING", (0, 0), (-1, -1), 8),
            ]
        )
    )
    story.extend([metric_table, Spacer(1, 18), Paragraph("Pending Issues", styles["Heading2"])])

    rows = [["Issue", "Status", "Title"]]
    rows.extend([issue.key, issue.status, Paragraph(issue.title, styles["BodyText"])] for issue in issues)
    issue_table = Table(rows, colWidths=[78, 76, 360], repeatRows=1)
    issue_table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#e5e7eb")),
                ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#d1d5db")),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
                ("TOPPADDING", (0, 0), (-1, -1), 6),
            ]
        )
    )
    story.append(issue_table)
    doc.build(story)


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("source", type=Path, help="Tracker source (.docx or markdown)")
    parser.add_argument("--output", "-o", type=Path, default=None, help="Output PDF path")
    parser.add_argument("--print-counts", action="store_true", help="Print derived status counts")
    args = parser.parse_args()

    issues = parse_issues(read_source(args.source))
    counts = status_counts(issues)
    if args.print_counts:
        print(f"PARTIAL={counts['PARTIAL']} OPEN={counts['OPEN']} TOTAL={counts['TOTAL']}")
    output = args.output or args.source.with_suffix(".pdf")
    build_pdf(args.source, output)


if __name__ == "__main__":
    main()
