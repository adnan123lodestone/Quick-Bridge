from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUTPUT = Path(__file__).with_name("QuickBridge_Meta_Integration_QA_Test_Plan.docx")

BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
INK = "0B2545"
MUTED = "5B6778"
TABLE_FILL = "E8EEF5"
LIGHT_FILL = "F4F6F9"
RISK_FILL = "FCEEEE"
SUCCESS_FILL = "EAF6EF"
GOLD_FILL = "FFF7DB"
BORDER = "B7C4D6"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin_name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin_name}"))
        if node is None:
            node = OxmlElement(f"w:{margin_name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths_dxa: list[int], indent_dxa: int = 120) -> None:
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    tbl = table._tbl
    tbl_pr = tbl.tblPr

    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:type"), "dxa")
    tbl_w.set(qn("w:w"), str(sum(widths_dxa)))

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:type"), "dxa")
    tbl_ind.set(qn("w:w"), str(indent_dxa))

    tbl_layout = tbl_pr.find(qn("w:tblLayout"))
    if tbl_layout is None:
        tbl_layout = OxmlElement("w:tblLayout")
        tbl_pr.append(tbl_layout)
    tbl_layout.set(qn("w:type"), "fixed")

    old_grid = tbl.tblGrid
    if old_grid is not None:
        tbl.remove(old_grid)
    grid = OxmlElement("w:tblGrid")
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    tbl.insert(0, grid)

    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            cell.width = Pt(widths_dxa[idx] / 20)
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:type"), "dxa")
            tc_w.set(qn("w:w"), str(widths_dxa[idx]))
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_margins(cell)


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_paragraph_keep_with_next(paragraph) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    keep_next = p_pr.find(qn("w:keepNext"))
    if keep_next is None:
        keep_next = OxmlElement("w:keepNext")
        p_pr.append(keep_next)


def set_run_style(run, bold=False, color=None, size=None) -> None:
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    if size:
        run.font.size = Pt(size)


def style_paragraph(paragraph, before=0, after=6, line=1.25) -> None:
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line


def add_heading(doc: Document, text: str, level: int = 1):
    p = doc.add_paragraph()
    if level == 1:
        style_paragraph(p, before=18, after=10, line=1.25)
        run = p.add_run(text)
        set_run_style(run, bold=True, color=BLUE, size=16)
    elif level == 2:
        style_paragraph(p, before=14, after=7, line=1.25)
        run = p.add_run(text)
        set_run_style(run, bold=True, color=BLUE, size=13)
    else:
        style_paragraph(p, before=10, after=5, line=1.25)
        run = p.add_run(text)
        set_run_style(run, bold=True, color=DARK_BLUE, size=12)
    set_paragraph_keep_with_next(p)
    return p


def add_body(doc: Document, text: str):
    p = doc.add_paragraph()
    style_paragraph(p)
    run = p.add_run(text)
    set_run_style(run, color=INK, size=11)
    return p


def add_bullet(doc: Document, text: str):
    p = doc.add_paragraph(style="List Bullet")
    style_paragraph(p, after=4)
    run = p.add_run(text)
    set_run_style(run, color=INK, size=11)
    return p


def add_number(doc: Document, text: str):
    p = doc.add_paragraph(style="List Number")
    style_paragraph(p, after=4)
    run = p.add_run(text)
    set_run_style(run, color=INK, size=11)
    return p


def add_callout(doc: Document, label: str, text: str, fill: str = LIGHT_FILL):
    table = doc.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    set_table_geometry(table, [9360])
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    p = cell.paragraphs[0]
    style_paragraph(p, after=2)
    r1 = p.add_run(label + ": ")
    set_run_style(r1, bold=True, color=DARK_BLUE, size=10.5)
    r2 = p.add_run(text)
    set_run_style(r2, color=INK, size=10.5)
    doc.add_paragraph()
    return table


def add_table(doc: Document, headers: list[str], rows: list[list[str]], widths_dxa: list[int], header_fill: str = TABLE_FILL):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    set_table_geometry(table, widths_dxa)
    header_cells = table.rows[0].cells
    for idx, header in enumerate(headers):
        set_cell_shading(header_cells[idx], header_fill)
        p = header_cells[idx].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        style_paragraph(p, after=0, line=1.15)
        run = p.add_run(header)
        set_run_style(run, bold=True, color=DARK_BLUE, size=9.5)
    set_repeat_table_header(table.rows[0])
    for row in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(row):
            p = cells[idx].paragraphs[0]
            style_paragraph(p, after=0, line=1.15)
            run = p.add_run(value)
            set_run_style(run, color=INK, size=9.5)
    set_table_geometry(table, widths_dxa)
    doc.add_paragraph()
    return table


def add_test_case(doc: Document, case_id: str, title: str, objective: str, data: str, steps: list[str], expected: list[str], evidence: str) -> None:
    add_heading(doc, f"{case_id} - {title}", 3)
    add_table(
        doc,
        ["Field", "Detail"],
        [
            ["Objective", objective],
            ["Test Data", data],
            ["Evidence", evidence],
        ],
        [1700, 7660],
        header_fill=LIGHT_FILL,
    )
    add_body(doc, "Steps")
    for step in steps:
        add_number(doc, step)
    add_body(doc, "Expected Results")
    for item in expected:
        add_bullet(doc, item)


def configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    for name in ("List Bullet", "List Number"):
        style = styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
        style.font.size = Pt(11)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.line_spacing = 1.25
        style.paragraph_format.left_indent = Inches(0.375)
        style.paragraph_format.first_line_indent = Inches(-0.188)

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    style_paragraph(header, after=0, line=1)
    run = header.add_run("QuickBridge Meta QA")
    set_run_style(run, color=MUTED, size=9)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    style_paragraph(footer, after=0, line=1)
    run = footer.add_run("Confidential QA test plan - do not store Meta secrets in this document")
    set_run_style(run, color=MUTED, size=8.5)


def add_title_page(doc: Document) -> None:
    p = doc.add_paragraph()
    style_paragraph(p, before=0, after=3, line=1.15)
    run = p.add_run("QuickBridge Meta Integration QA Test Plan")
    set_run_style(run, bold=True, color=INK, size=24)

    p = doc.add_paragraph()
    style_paragraph(p, after=14, line=1.15)
    run = p.add_run("Facebook and Instagram connector validation for Salesforce QA")
    set_run_style(run, color=MUTED, size=13)

    add_table(
        doc,
        ["Document Field", "Value"],
        [
            ["Connector", "Meta / Facebook / Instagram"],
            ["Repository", "QuickBridge Salesforce"],
            ["Audience", "QA analysts, Salesforce admins, implementation consultants"],
            ["Scope", "Connector catalog tile, configuration records, webhook verification, lead ads, comment attribution, work routing, scheduler routing, and negative-path validation"],
            ["Out of Scope", "Live Meta Marketing API polling for campaign and insights data unless that Phase 2 API implementation has been added to the target org"],
        ],
        [2100, 7260],
    )
    add_callout(
        doc,
        "Phase 1 scope",
        "This plan validates the Meta connector foundation currently in the repo. CampaignSync, InsightsSync, and Reconciliation work items are routed and usage-counted, but the live Meta Marketing API polling layer is not treated as complete in this document.",
        GOLD_FILL,
    )


def add_overview(doc: Document) -> None:
    add_heading(doc, "1. QA Scope And Entry Criteria", 1)
    add_body(
        doc,
        "Use this document to validate that the Meta connector can be discovered, configured, secured, and exercised through Salesforce-only and live Meta-assisted test paths. The plan is intentionally split so QA can run a sandbox smoke test without live Meta credentials, then add live webhook and lead-form checks when Meta app assets are available.",
    )
    add_heading(doc, "Entry Criteria", 2)
    for item in [
        "The Meta implementation metadata and Apex are deployed to the QA org.",
        "The Meta static resource is deployed and visible as /resource/Meta_Logo.",
        "QuickBridge config panel access is available to the QA user.",
        "A QA admin can run anonymous Apex scripts or a deployment engineer has run the seed scripts.",
        "No real app secret, token, or customer PII is copied into this document or attached to defect tickets.",
    ]:
        add_bullet(doc, item)
    add_heading(doc, "Live Meta Prerequisites", 2)
    for item in [
        "A Meta app configured for Facebook Page and Instagram webhook testing.",
        "A Facebook Page ID, Instagram Business Account ID, Ad Account ID, and test lead form where live checks are required.",
        "A webhook callback URL that maps to the org endpoint /services/apexrest/meta-webhook/ and can be reached by Meta.",
        "Secret reference values created in the org's approved secret storage approach. The connector config stores references, not plaintext secrets.",
    ]:
        add_bullet(doc, item)


def add_traceability(doc: Document) -> None:
    add_heading(doc, "2. Traceability Matrix", 1)
    add_table(
        doc,
        ["Area", "Primary Repo Assets", "QA Focus"],
        [
            ["Catalog tile", "Integration_Product_Default.Meta, Integration_Connector.Meta, Meta_Logo static resource, quickbridgeConfigPanel LWC", "Tile appears as Facebook / Instagram with a working logo and setup action."],
            ["Configuration", "Connector_Config_Field.Meta_*, scripts/setupMetaConnector.apex, Connector_Instance__c, Connector_Config_Value__c", "Required values exist, placeholders are replaced before live traffic, and secret values are references."],
            ["Webhook security", "MetaWebhookController, MetaWebhookVerifier, MetaConfigService", "GET challenge succeeds only with the configured verify token; POST accepts only valid X-Hub-Signature-256 payloads."],
            ["Lead Ads", "MetaPayloadNormalizer, MetaLeadAdsProcessor, MetaLeadService, MetaCampaignMemberService", "Leadgen payloads create or update Leads, Social Interactions, and Campaign Members idempotently."],
            ["Comment attribution", "MetaCommentAttributionProcessor, MetaSocialInteractionService, MetaAttributionService", "Facebook and Instagram comments resolve to campaign/ad data and create Leads only when the configured lead creation mode allows it."],
            ["Work routing", "Integration_Connector_Capability.Meta_*, MetaIntegrationWorker, IntegrationWorkRouter", "LeadAd, CommentAttribution, WebhookEvent, CampaignSync, InsightsSync, and Reconciliation operations route safely."],
            ["Scheduling", "MetaScheduledOperationAdapter, Integration_Work_Setting.Meta_Default", "Scheduled Campaign and Insights jobs enqueue Meta work items without using unapproved adapters."],
            ["Object mapping", "Integration_Object_Map.Meta_Campaign, Integration_Object_Map.Meta_Ad_Set, Integration_Object_Map.Meta_Ad", "Campaign, Meta Ad Set, and Meta Ad data can support attribution lookups."],
        ],
        [1500, 4250, 3610],
    )


def add_test_data(doc: Document) -> None:
    add_heading(doc, "3. Test Data And Setup", 1)
    add_heading(doc, "3.1 Seed And Configuration Steps", 2)
    for step in [
        "Deploy or confirm the Meta source package in the QA org.",
        "Run scripts/reseedIntegrationProducts.apex after metadata deployment to refresh runtime product catalog rows.",
        "Run scripts/setupMetaConnector.apex to create or update the Meta Connector_Instance__c and Connector_Config_Value__c rows.",
        "Replace every REPLACE_WITH_* placeholder before enabling live Meta traffic.",
        "Confirm the three secret-reference fields are marked Is_Secret_Reference__c = true: appSecretRef, accessTokenRef, webhookVerifyTokenRef.",
    ]:
        add_number(doc, step)
    add_table(
        doc,
        ["Field Key", "Example QA Value", "Notes"],
        [
            ["appId", "QA Meta App ID", "Non-secret identifier."],
            ["appSecretRef", "QB_META_APP_SECRET_QA", "Reference only. Do not store the plaintext app secret in Salesforce config text."],
            ["accessTokenRef", "QB_META_ACCESS_TOKEN_QA", "Reference only. Must point to an approved secret store or Named Credential strategy."],
            ["webhookVerifyTokenRef", "QB_META_VERIFY_TOKEN_QA", "Reference or token value used by webhook challenge validation."],
            ["adAccountId", "act_<qa_ad_account_id>", "Required for live campaign and reporting checks."],
            ["pageId", "<qa_facebook_page_id>", "Required for live Page webhook checks."],
            ["instagramBusinessAccountId", "<qa_ig_business_account_id>", "Required for Instagram comment checks."],
            ["defaultLeadSource", "Meta Lead Ad", "Expected fallback LeadSource."],
            ["defaultCampaignMemberStatus", "Responded", "Must be valid for the Campaign Member Status list in the org."],
            ["leadCreationMode", "Manual Review, Keyword Match, or Always", "Controls whether comments create Leads."],
            ["keywordRulesJson", "[\"pricing\", \"quote\"]", "Only used for Keyword Match mode."],
            ["defaultFallbackCampaignId", "<optional Campaign Id>", "Use only when QA intentionally validates fallback attribution."],
        ],
        [2300, 3000, 4060],
    )

    add_heading(doc, "3.2 SOQL Verification Snippets", 2)
    add_body(doc, "Use these snippets from Developer Console, Workbench, or sf data query. Replace the target org alias as appropriate.")
    snippets = [
        "SELECT Label, Connector_Key__c, Product_Key__c, Logo_URL__c, Status__c FROM Integration_Product__c WHERE Product_Key__c = 'meta'",
        "SELECT Id, Name, Active__c, Health_Status__c, Entitlement_Status__c FROM Connector_Instance__c WHERE Connector_Key__c = 'meta'",
        "SELECT Field_Key__c, Value__c, Is_Secret_Reference__c FROM Connector_Config_Value__c WHERE Connector_Key__c = 'meta' ORDER BY Field_Key__c",
        "SELECT Integration__c, Operation__c, Source_Object__c, Source_Record_Id__c, Status__c FROM Integration_Work_Item__c WHERE Integration__c = 'Meta' ORDER BY CreatedDate DESC",
    ]
    for snippet in snippets:
        add_callout(doc, "SOQL", snippet, LIGHT_FILL)

    add_heading(doc, "3.3 Sample Webhook Payloads", 2)
    add_callout(
        doc,
        "Lead Ads payload",
        '{"object":"page","entry":[{"id":"page-1","changes":[{"field":"leadgen","value":{"leadgen_id":"qa-leadgen-001","ad_id":"qa-ad-001","adset_id":"qa-adset-001","campaign_id":"qa-campaign-001","form_id":"qa-form-001","page_id":"page-1"}}]}]}',
        LIGHT_FILL,
    )
    add_callout(
        doc,
        "Instagram comment payload",
        '{"object":"instagram","entry":[{"id":"ig-1","changes":[{"field":"comments","value":{"id":"qa-comment-001","media_id":"qa-media-001","from":{"id":"qa-profile-001","username":"qa_buyer"},"text":"Need pricing"}}]}]}',
        LIGHT_FILL,
    )


def add_test_cases(doc: Document) -> None:
    add_heading(doc, "4. Test Cases", 1)
    cases = [
        (
            "META-QA-001",
            "Product Catalog Metadata And Runtime Product",
            "Confirm the Facebook / Instagram catalog row exists and points to the Meta static resource.",
            "Integration_Product_Default.Meta and reseeded Integration_Product__c row.",
            [
                "Run the Integration_Product__c SOQL query from section 3.2.",
                "Confirm Product_Key__c and Connector_Key__c both equal meta.",
                "Confirm Product_Label__c or Label displays Facebook / Instagram.",
                "Confirm Logo_URL__c equals /resource/Meta_Logo.",
                "Confirm Status__c is Available before activation and changes appropriately after subscription setup if the org uses subscription state.",
            ],
            [
                "A single Meta product row is returned.",
                "The row uses the Meta logo resource and does not reference a broken external image.",
                "No duplicate Facebook-only or Instagram-only product rows appear unless explicitly added as future separate products.",
            ],
            "Screenshot of the query result and config panel product tile.",
        ),
        (
            "META-QA-002",
            "Config Panel Tile Visibility",
            "Verify QA can see the Meta tile in the QuickBridge config panel.",
            "QA user with the same access used for existing QuickBooks, Shopify, Stripe, Authorize.Net, FedEx, DHL, PayPal, UPS, and Klaviyo tiles.",
            [
                "Open the QuickBridge config panel in the QA org.",
                "Refresh the page with cache disabled if the component was recently deployed.",
                "Review both Subscribed Products and Available Products sections.",
                "Use browser dev tools Network tab only if the tile is still missing, and confirm the product service response contains meta.",
                "Click Setup Now or the active tile action, depending on subscription state.",
            ],
            [
                "A Facebook / Instagram tile appears with the Meta logo.",
                "The tile is not hidden behind the Available Products toggle.",
                "Clicking the tile opens the same configuration workflow pattern as other metadata-driven products.",
            ],
            "Screenshot of the full product grid showing the Meta tile.",
        ),
        (
            "META-QA-003",
            "Connector Instance And Config Values",
            "Confirm setupMetaConnector.apex creates the runtime records needed by the connector.",
            "Run scripts/setupMetaConnector.apex in the QA org.",
            [
                "Run the Connector_Instance__c SOQL query from section 3.2.",
                "Run the Connector_Config_Value__c SOQL query from section 3.2.",
                "Count the returned config rows.",
                "Inspect values for REPLACE_WITH_* placeholders.",
                "Verify appSecretRef, accessTokenRef, and webhookVerifyTokenRef are marked as secret references.",
            ],
            [
                "One active Connector_Instance__c exists for Connector_Key__c = meta.",
                "Twenty two config value rows exist after the setup script.",
                "Live testing does not proceed while any required placeholder value remains.",
                "Secret values are not visible as plaintext application secrets or access tokens.",
            ],
            "SOQL export or screenshot with sensitive values redacted.",
        ),
        (
            "META-QA-004",
            "Editable Config Fields",
            "Validate that the UI exposes the intended Meta configuration fields.",
            "Integration_Connector.Meta editable field list and Connector_Config_Field.Meta_* metadata.",
            [
                "Open the Meta configuration page from the tile.",
                "Verify fields for app, account, page, Instagram account, webhook, Lead Ads, comment attribution, insights, sync frequency, lead source, campaign status, keyword rules, attribution window, and fallback campaign.",
                "Update a non-secret field such as defaultLeadSource or syncFrequency.",
                "Save and re-query Connector_Config_Value__c.",
                "Revert the field to the approved QA value.",
            ],
            [
                "Only Meta-related editable fields are shown.",
                "Saved non-secret values persist to Connector_Config_Value__c.",
                "Secret reference fields behave as references and do not reveal plaintext secrets back to the browser.",
            ],
            "Before/after screenshots of the config panel and SOQL verification.",
        ),
        (
            "META-QA-005",
            "Webhook Verification Challenge",
            "Validate that Meta's GET verification flow accepts only the configured token.",
            "A QA webhook verify token and accessible /services/apexrest/meta-webhook/ endpoint.",
            [
                "Set webhookVerifyTokenRef to the QA-approved verify token or reference used by the implementation.",
                "Send a GET request with hub.mode=subscribe, hub.verify_token=<valid token>, and hub.challenge=qa-challenge-001.",
                "Repeat the GET request with an invalid token.",
                "Review response body and HTTP status for both calls.",
            ],
            [
                "The valid request returns qa-challenge-001 with HTTP 200.",
                "The invalid request returns Forbidden with HTTP 403.",
                "No sensitive config values are included in either response.",
            ],
            "HTTP request/response transcript with tokens redacted.",
        ),
        (
            "META-QA-006",
            "Webhook POST Signature Enforcement",
            "Confirm incoming POST payloads require a valid X-Hub-Signature-256 header.",
            "Sample payload from section 3.3 and HMAC SHA-256 signature generated with the QA app secret.",
            [
                "Generate a sha256=<hex hmac> signature for the exact payload body using the QA app secret.",
                "POST the payload to /services/apexrest/meta-webhook/ with X-Hub-Signature-256 set to the valid signature.",
                "Repeat the POST with a malformed or mismatched signature.",
                "Repeat the POST with an empty body.",
            ],
            [
                "Valid signed payload returns success true, message Meta webhook accepted, and a workItemId.",
                "Invalid signature returns HTTP 401 and success false.",
                "Empty body returns HTTP 400 and success false.",
                "The accepted request creates an Integration_Work_Item__c with Integration__c = Meta and Operation__c = WebhookEvent.",
            ],
            "Request transcript plus Integration_Work_Item__c query result.",
        ),
        (
            "META-QA-007",
            "Lead Ads Webhook Processing",
            "Verify a Lead Ads webhook event creates the expected Salesforce records.",
            "A Campaign with Meta_Campaign_Id__c matching the test campaign ID; optional Meta_Ad_Set__c and Meta_Ad__c records for stronger attribution.",
            [
                "Create or confirm a test Campaign with Meta_Campaign_Id__c = qa-campaign-001 and IsActive = true.",
                "Create optional Meta_Ad_Set__c and Meta_Ad__c records with IDs matching qa-adset-001 and qa-ad-001.",
                "Submit a valid signed Lead Ads webhook payload.",
                "Run the Integration_Work_Item__c dispatcher if the org does not auto-process queued work immediately.",
                "Query Lead, Social_Interaction__c, and CampaignMember records for qa-leadgen-001.",
            ],
            [
                "A Lead exists with Meta_Leadgen_Id__c = qa-leadgen-001.",
                "Lead source defaults to Meta Lead Ad unless QA configured a different approved value.",
                "A Social_Interaction__c exists and links to the Lead.",
                "A CampaignMember exists when autoCreateCampaignMember is true and attribution resolves to a Campaign.",
            ],
            "Record IDs for Lead, Social Interaction, Campaign, and Campaign Member.",
        ),
        (
            "META-QA-008",
            "Lead Ads Idempotency",
            "Confirm duplicate leadgen processing does not create duplicate Leads, Social Interactions, or Campaign Members.",
            "Same payload and records used for META-QA-007.",
            [
                "Submit the same valid signed Lead Ads payload a second time.",
                "Process the queued work item if needed.",
                "Query Lead where Meta_Leadgen_Id__c = qa-leadgen-001.",
                "Query Social_Interaction__c and CampaignMember for the same Lead and Campaign.",
            ],
            [
                "Only one Lead exists for the leadgen ID.",
                "Only one Social_Interaction__c exists for the leadgen ID.",
                "Only one CampaignMember exists for the Lead and Campaign.",
                "The duplicate work item does not produce an unhandled exception.",
            ],
            "Aggregate query counts before and after the duplicate submission.",
        ),
        (
            "META-QA-009",
            "Instagram Comment Attribution - Manual Review",
            "Verify comment events are captured without automatic Lead creation when leadCreationMode is Manual Review.",
            "leadCreationMode = Manual Review; Meta_Ad__c with Instagram_Media_Id__c = qa-media-001.",
            [
                "Set leadCreationMode to Manual Review.",
                "Create a Campaign and Meta_Ad__c linked to Instagram_Media_Id__c = qa-media-001.",
                "Submit a valid signed Instagram comment payload from section 3.3.",
                "Process queued work if required.",
                "Query Social_Interaction__c where Meta_Comment_Id__c = qa-comment-001.",
            ],
            [
                "A Social_Interaction__c is created with the comment text and commenter identity where available.",
                "The interaction resolves to the expected Campaign and/or Ad.",
                "Lead__c remains blank in Manual Review mode.",
            ],
            "Social Interaction screenshot or SOQL result.",
        ),
        (
            "META-QA-010",
            "Instagram Comment Attribution - Keyword Match",
            "Verify comment events create Leads only when configured keyword rules match.",
            "leadCreationMode = Keyword Match; keywordRulesJson includes pricing.",
            [
                "Set leadCreationMode to Keyword Match and keywordRulesJson to [\"pricing\"].",
                "Submit a signed Instagram comment payload whose text includes pricing.",
                "Submit a second signed payload whose text does not include pricing.",
                "Process queued work if required.",
                "Query Social_Interaction__c and Lead records for both comments.",
            ],
            [
                "The matching comment creates or links a Lead and can create a CampaignMember when attribution and config allow it.",
                "The non-matching comment still creates a Social Interaction but does not create a Lead.",
                "Keyword matching is case-insensitive.",
            ],
            "SOQL results for both comment IDs.",
        ),
        (
            "META-QA-011",
            "Attribution Fallback Paths",
            "Validate campaign, ad set, ad, post, and media identifiers resolve attribution consistently.",
            "Campaign, Meta_Ad_Set__c, and Meta_Ad__c rows with matching Meta IDs.",
            [
                "Prepare one Campaign with Meta_Campaign_Id__c matching a test payload campaign_id.",
                "Prepare one Meta_Ad_Set__c with Meta_Ad_Set_Id__c matching a test payload adset_id.",
                "Prepare one Meta_Ad__c with Meta_Ad_Id__c, Post_Id__c, and Instagram_Media_Id__c matching test payload identifiers.",
                "Run Lead Ads and comment payloads that include each identifier path.",
                "Inspect Social_Interaction__c Attribution_Status__c and related lookup fields.",
            ],
            [
                "Attribution resolves to the most specific available ad/ad set/campaign record.",
                "Missing campaign data results in Needs Review or Unattributed rather than a silent failure.",
                "Fallback campaign is used only when explicitly configured and intended.",
            ],
            "Matrix of payload identifier used versus resolved Salesforce record.",
        ),
        (
            "META-QA-012",
            "Scheduler And Work Routing",
            "Confirm scheduled Meta jobs enqueue the correct work operations.",
            "Scheduled_Integration_Job__c records for Source_Object__c = Campaign and Insights.",
            [
                "Create or identify a Scheduled_Integration_Job__c with Integration_Type__c = Meta and Source_Object__c = Campaign.",
                "Execute the scheduled dispatcher path used by the org, or run the adapter through a controlled QA Apex script.",
                "Repeat for Source_Object__c = Insights.",
                "Query Integration_Work_Item__c for Integration__c = Meta.",
            ],
            [
                "Campaign source object enqueues Operation__c = CampaignSync.",
                "Insights source object enqueues Operation__c = InsightsSync.",
                "The configured allowlist permits MetaScheduledOperationAdapter and MetaIntegrationWorker only as expected.",
                "No live Marketing API data import is expected unless Phase 2 polling has been implemented.",
            ],
            "Integration_Work_Item__c query results for both operations.",
        ),
        (
            "META-QA-013",
            "Unsupported Or Malformed Work",
            "Validate negative-path behavior and error logging.",
            "Controlled QA Apex execution of MetaIntegrationWorker with malformed input.",
            [
                "Submit an empty webhook POST and confirm HTTP 400.",
                "Submit a bad signature and confirm HTTP 401.",
                "Queue or invoke a Meta work item with Operation__c = Unsupported.",
                "Queue or invoke a Meta work item with invalid JSON payload.",
                "Review error logging according to the org's ErrorLogUtility pattern.",
            ],
            [
                "Bad webhook requests fail before work is created.",
                "Unsupported operations raise a useful MetaIntegrationWorkerException.",
                "Unexpected processing failures are logged with Meta context.",
                "No partial sensitive payload or token is exposed in user-facing error messages.",
            ],
            "Error log IDs and sanitized request details.",
        ),
        (
            "META-QA-014",
            "Regression Check For Existing Products",
            "Verify the Meta catalog addition did not break existing product tiles.",
            "Existing QuickBridge product catalog rows.",
            [
                "Open the config panel after reseeding products.",
                "Confirm existing subscribed and available products still render.",
                "Specifically confirm UPS logo renders after the recent static resource correction.",
                "Open at least one existing product configuration panel and save a harmless non-secret field only if QA policy allows.",
            ],
            [
                "Existing product tiles remain visible.",
                "Existing logos render without broken image placeholders.",
                "No Meta-specific logic changes labels, state, or behavior for non-Meta products.",
            ],
            "Screenshot of product grid after Meta is visible.",
        ),
    ]
    for case in cases:
        add_test_case(doc, *case)


def add_execution_checklist(doc: Document) -> None:
    add_heading(doc, "5. QA Execution Checklist", 1)
    add_table(
        doc,
        ["Status", "Checkpoint", "Owner / Evidence"],
        [
            ["[ ]", "Deployment and seed scripts completed in QA.", ""],
            ["[ ]", "Meta tile appears with /resource/Meta_Logo.", ""],
            ["[ ]", "Config values are present and required placeholders are replaced for live tests.", ""],
            ["[ ]", "Secret fields are references and not plaintext secrets.", ""],
            ["[ ]", "Webhook GET challenge valid and invalid paths tested.", ""],
            ["[ ]", "Webhook POST valid signature, invalid signature, and empty body tested.", ""],
            ["[ ]", "Lead Ads record creation and idempotency tested.", ""],
            ["[ ]", "Instagram/Facebook comment attribution tested for Manual Review and Keyword Match modes.", ""],
            ["[ ]", "Scheduler routing tested for CampaignSync and InsightsSync.", ""],
            ["[ ]", "Regression pass completed for existing product tiles.", ""],
            ["[ ]", "Known exclusions acknowledged by QA lead.", ""],
            ["[ ]", "Defects logged with sanitized evidence and no secrets.", ""],
        ],
        [900, 5700, 2760],
    )
    add_heading(doc, "Defect Template", 2)
    add_table(
        doc,
        ["Field", "Required Detail"],
        [
            ["Title", "Short failure summary with test case ID."],
            ["Environment", "Org alias, sandbox name, build/deploy job if known."],
            ["Steps", "Exact steps, payload type, and whether live Meta or Salesforce-only execution was used."],
            ["Expected", "Expected result from this test plan."],
            ["Actual", "Observed response, record state, log entry, or screenshot."],
            ["Evidence", "Sanitized screenshots, SOQL output, request ID, work item ID, and log ID."],
            ["Security Check", "Confirm no secrets, tokens, customer PII, or real lead data are attached."],
        ],
        [1700, 7660],
    )


def add_signoff(doc: Document) -> None:
    add_heading(doc, "6. Exit Criteria And Sign-Off", 1)
    add_body(
        doc,
        "Meta Phase 1 QA is considered complete when all in-scope test cases pass, known Phase 2 exclusions are accepted, and all defects are either resolved or dispositioned for a later release.",
    )
    add_table(
        doc,
        ["Role", "Name", "Date", "Decision / Notes"],
        [
            ["QA Lead", "", "", ""],
            ["Salesforce Lead", "", "", ""],
            ["Implementation Owner", "", "", ""],
            ["Client/Product Owner", "", "", ""],
        ],
        [1900, 2200, 1400, 3860],
    )


def main() -> None:
    doc = Document()
    configure_document(doc)
    add_title_page(doc)
    add_overview(doc)
    add_traceability(doc)
    add_test_data(doc)
    add_test_cases(doc)
    add_execution_checklist(doc)
    add_signoff(doc)

    core_props = doc.core_properties
    core_props.title = "QuickBridge Meta Integration QA Test Plan"
    core_props.subject = "QA validation steps for the QuickBridge Meta Facebook and Instagram connector"
    core_props.keywords = "QuickBridge, Salesforce, Meta, Facebook, Instagram, QA, test plan"
    core_props.comments = "Generated from repository metadata and Apex implementation. No secrets included."

    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
